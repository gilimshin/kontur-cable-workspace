"""Выгрузка результата в .xlsx в формате вкладки «кВ» (формулы как в рабочем макросе)."""
import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side


def _sum_formula(parts):
    parts = [p for p in parts if p]
    if not parts:
        return 0
    if len(parts) == 1:
        return parts[0]
    return '=' + '+'.join(f'{p:g}' for p in parts)


def build_xlsx(result, rows=None):
    rows = rows if rows is not None else result['table']
    k_g = result['params']['k_ground']
    k_p = result['params']['k_pipe']
    wb = Workbook()
    ws = wb.active
    ws.title = 'кВ'
    thin = Side(style='thin', color='999999')
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    head = PatternFill('solid', fgColor='DDEBF7')
    hdr = ['Каб. Жур', 'Кабель', 'Сечение', 'в земле', 'в ГНБ (профиль)', 'в трубе ', 'Заводка', 'Итоговая  длина', 'Округлен.']
    for i, h in enumerate(hdr, 1):
        c = ws.cell(1, i, h)
        c.font = Font(bold=True)
        c.fill = head
        c.border = border
        c.alignment = Alignment(wrap_text=True, vertical='center')
    r = 2
    for row in rows:
        ws.cell(r, 1, row['name'])
        ws.cell(r, 2, row['mark'])
        ws.cell(r, 3, row['section'])
        ws.cell(r, 4, _sum_formula(row['ground_parts']))
        ws.cell(r, 5, _sum_formula(row['gnb_parts']))
        ws.cell(r, 6, _sum_formula(row['pipe_parts']))
        ws.cell(r, 7, f"={row['zav_from']:g}+{row['zav_to']:g}")
        ws.cell(r, 8, f'=D{r}*{k_g:g}+E{r}+F{r}*{k_p:g}+G{r}')
        ws.cell(r, 9, f'=ROUND(H{r},0)')
        for c in range(1, 10):
            ws.cell(r, c).border = border
        r += 1
    last = r - 1
    ws.cell(r, 4, f'=SUM(D2:D{last})')
    ws.cell(r, 7, f'=SUM(G2:G{last})')
    ws.cell(r, 9, f'=SUM(I2:I{last})').font = Font(bold=True)

    # таблицы труб (длина × количество), как в рабочей вкладке
    r += 4
    col0 = 1
    pipes = result['summary']['pipes']
    for name, block in pipes.items():
        ws.cell(r, col0, name).font = Font(bold=True)
        ws.cell(r, col0 + 1, 'Длина')
        ws.cell(r, col0 + 2, 'Кол-во')
        ws.cell(r, col0 + 3, 'Общая длина')
        rr = r + 1
        for p in block['rows']:
            ws.cell(rr, col0, p['span'])
            ws.cell(rr, col0 + 1, p['len'])
            ws.cell(rr, col0 + 2, p['count'])
            L = ws.cell(rr, col0 + 1).column_letter
            Cn = ws.cell(rr, col0 + 2).column_letter
            ws.cell(rr, col0 + 3, f'={L}{rr}*{Cn}{rr}')
            rr += 1
        D = ws.cell(rr, col0 + 3).column_letter
        ws.cell(rr, col0 + 3, f'=SUM({D}{r + 1}:{D}{rr - 1})').font = Font(bold=True)
        col0 += 5
    r = max(r + 2 + max((len(b['rows']) for b in pipes.values()), default=0), r + 2) + 2

    s = result['summary']
    ws.cell(r, 1, 'Общий').font = Font(bold=True)
    ws.cell(r, 2, 'Исп. чертеж')
    items = [('L без ГНБ (плн.)', s['trench_no_gnb']), ('Траншея без труб', s['trench_ground']),
             ('Трубы (длина трассы в трубах)', s['trench_pipe']), ('ГНБ (профиль)', s['gnb_prof'])]
    items += [(k, v['total']) for k, v in pipes.items()]
    items += [('Кол-во КЛ', s['cables_count'])]
    for i, (k, v) in enumerate(items, 1):
        ws.cell(r + i, 1, k)
        ws.cell(r + i, 2, v)

    # лист с разбивкой маршрутов
    w2 = wb.create_sheet('Маршруты')
    w2.append(['КЛ', 'Спецификация, м', 'По маршруту, м', 'Статус', 'Маршрут (точки)', 'Участки'])
    for c in result['cables']:
        steps = ' | '.join(f"т.{x['frm']}–т.{x['to']}: {x['len']:g} ({ {'ground': 'земля', 'pipe': 'труба', 'gnb': 'ГНБ'}[x['kind']] })"
                           for x in c.get('steps', []))
        w2.append([f"{c['count']}КЛ {c['frm']} (т.{c['a']}) - {c['to']} (т.{c['b']})", c['spec_len'], c.get('route_len'),
                   c.get('status'), ' → '.join(c.get('path', [])), steps])
    for ws_ in (ws, w2):
        for col, w in zip('ABCDEFGHI', (34, 20, 14, 14, 14, 14, 10, 14, 10)):
            ws_.column_dimensions[col].width = w
    w2.column_dimensions['E'].width = 60
    w2.column_dimensions['F'].width = 120
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
